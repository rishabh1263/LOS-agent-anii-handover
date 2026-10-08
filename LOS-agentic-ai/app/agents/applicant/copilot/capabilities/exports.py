"""
DOWNLOADS (MASTER SPEC sections 15.3 / 15.5; config app/config/product_flow.yaml `exports`).

Excel (.xlsx), Doc (.docx) and PDF for ONE case (a summary) or a FILTERED LIST of the caller's own cases.

    ONE CONTENT     `case_content` / `list_content` build the same title, sections and tables once; three writers
                    only lay them out (no channel has its own facts -- section 19).
    LIBRARIES       openpyxl (installed) for Excel; PyMuPDF (installed, the handoff note's writer) for PDF; a .docx
                    is plain Office Open XML written with the standard library's zipfile -- no new dependency.
    SCOPE           a case: the ordinary ownership check; a list: the same list query as the chat and the home
                    table (only the caller's own cases). The route hands out a short-lived signed link bound to the
                    caller, re-checks scope when it is opened, and audits both.
    MASKING         every text is masked by the PII rule (sensitivity) and the abuse rule (abuse_guard) before it is
                    written -- an export never carries what the screen would not show.
"""

from __future__ import annotations

import base64
import hashlib
import hmac
import io
import json
import os
import threading
import time
import zipfile
from collections import defaultdict, deque
from datetime import datetime, timezone
from typing import Any
from xml.sax.saxutils import escape

FORMATS = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
           "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
           "pdf": "application/pdf"}

_LOCK = threading.Lock()
_CALLS: dict[str, deque] = defaultdict(deque)


def _cfg() -> dict[str, Any]:
    from app.agents.applicant.copilot.capabilities import product_flow

    return product_flow.cfg().get("exports") or {}


def _say(value: Any, lang: str, **values: Any) -> str:
    from app.agents.applicant.copilot.capabilities import product_flow

    return product_flow.say(value, lang, **values)


def _clean(value: Any) -> str:
    """The text as the screen may show it: PII masked, flagged words masked, no emojis."""
    from app.agents.applicant.copilot.answering import professional
    from app.agents.applicant.copilot.capabilities import abuse_guard
    from app.security import sensitivity

    text = "" if value is None else str(value)
    return professional.strip_emojis(abuse_guard.mask_text(sensitivity.mask_identifiers(text))).strip()


# --------------------------------------------------------------------------
# rate limit (section 17.6) and the signed link
# --------------------------------------------------------------------------

def allowed(subject: str, *, now: float | None = None) -> bool:
    now = now or time.time()
    with _LOCK:
        calls = _CALLS[subject]
        while calls and now - calls[0] > 60:
            calls.popleft()
        if len(calls) >= int(_cfg().get("rate_limit_per_minute", 10)):
            return False
        calls.append(now)
        return True


def reset() -> None:
    with _LOCK:
        _CALLS.clear()


def _key() -> bytes:
    """Derived from the existing data-encryption key (no new secret); a different purpose prefix than documents."""
    from app.store import crypto

    material = (os.getenv(crypto.ENV_KEY) or crypto._DEV_KEY).encode()
    return hashlib.sha256(b"los-export-v1:" + material).digest()


def _b64(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).decode().rstrip("=")


def sign(spec: dict[str, Any], subject: str, *, now: float | None = None) -> tuple[str, int]:
    ttl = int(_cfg().get("link_ttl_seconds", 300))
    payload = _b64(json.dumps({**spec, "s": subject, "e": int((now or time.time()) + ttl)},
                              separators=(",", ":")).encode())
    return f"{payload}.{_b64(hmac.new(_key(), payload.encode(), hashlib.sha256).digest())}", ttl


def verify(token: str, subject: str, *, now: float | None = None) -> dict[str, Any] | None:
    """The export spec when the link is genuine, unexpired and the SAME subject's; else None."""
    try:
        payload, mac = str(token or "").split(".", 1)
        good = _b64(hmac.new(_key(), payload.encode(), hashlib.sha256).digest())
        if not hmac.compare_digest(good, mac):
            return None
        data = json.loads(base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4)))
    except Exception:  # noqa: BLE001 - anything malformed is simply not a valid link
        return None
    if int(data.get("e", 0)) < int(now or time.time()) or data.get("s") != subject:
        return None
    return {k: v for k, v in data.items() if k not in ("s", "e")}


# --------------------------------------------------------------------------
# the content (one place; the writers only lay it out)
# --------------------------------------------------------------------------

def _today() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def case_content(case_id: str, claims: dict[str, Any], lang: str) -> dict[str, Any]:
    """{title, subtitle, sections: [{heading, columns, rows}]} for one case. Scope: the ordinary check."""
    from app.agents.applicant.copilot.answering import document_actions
    from app.agents.applicant.copilot.capabilities import case_form, product_flow, workspace
    from app.security.auth import get_subject
    from app.store import get_repository

    workspace.authorize(claims, case_id)
    repo = get_repository()
    application = repo.get_application(case_id)
    applicant = repo.get_applicant(application.applicant_id) if application else None
    c = _cfg()
    sections = c.get("sections") or {}
    from app.agents.applicant.copilot.answering import language_lock

    field_cols = language_lock.pick(c.get("field_columns"), lang) or ["Field", "Value"]
    stage = workspace._stage(case_id) or "--"
    status = product_flow.status_of(stage)
    head_rows = [[_column("case_id", lang), case_id],
                 [_column("app_id", lang), application.applicant_id if application else ""],
                 [_column("stage", lang), stage],
                 [_column("status", lang), product_flow.status_label(status, lang)
                  + f" ({product_flow.group_label(product_flow.group_of(status), lang)})"]]
    values = case_form.values_of(case_id)
    applicant_rows = [[case_form.label(f, lang), values.get(f)] for f in case_form.APPLICANT_FIELDS
                      if values.get(f) not in (None, "")]
    application_rows = [[case_form.label(f, lang), values.get(f)] for f in case_form.APPLICATION_FIELDS
                        if values.get(f) not in (None, "")]
    doc_cols = language_lock.pick(c.get("document_columns"), lang) or ["Document", "Party", "Status"]
    from app.agents.applicant.copilot.answering.answer import _readable

    docs = [[_readable(str(d.document_type)), str(d.party_role or "").replace("_", " ").title(),
             str(getattr(d.status, "value", d.status))] for d in repo.list_documents(case_id)]
    view = document_actions.build(case_id)
    pending = [str(r.get("label") or _readable(str(r.get("document_type") or r.get("slot") or "")))
               for key in ("reupload", "pending") for r in view.get(key) or [] if isinstance(r, dict)]
    out_sections = [
        {"heading": _say(sections.get("application"), lang), "columns": field_cols, "rows": head_rows
         + application_rows},
        {"heading": _say(sections.get("applicant"), lang), "columns": field_cols, "rows": applicant_rows},
        {"heading": _say(sections.get("documents"), lang), "columns": doc_cols, "rows": docs},
        {"heading": _say(sections.get("pending"), lang), "columns": [],
         "rows": [[p] for p in pending] or [[_say(c.get("nothing_pending"), lang)]]},
    ]
    user = str(get_subject(claims) or "")
    return _masked({"title": f"{_say((c.get('title') or {}).get('case'), lang)} -- {case_id}",
                    "subtitle": _say(c.get("generated"), lang, date=_today(), user=user),
                    "sections": out_sections, "file": _name("case", case_id)})


def list_content(claims: dict[str, Any], filters: dict[str, Any], lang: str) -> dict[str, Any]:
    """The filtered list of the caller's own cases (the home table's columns), capped at exports.max_rows."""
    from app.agents.applicant.copilot.answering import language_lock
    from app.agents.applicant.copilot.capabilities import case_list, product_flow
    from app.security.auth import get_subject

    c = _cfg()
    cap = int(c.get("max_rows", 500))
    query = case_list.ListQuery(**{k: v for k, v in filters.items() if k in product_flow.LIST_KEYS})
    rows: list[dict[str, Any]] = []
    page_size = int(case_list.cfg().get("max_page_size", 20))
    page = 0
    while len(rows) < cap:
        query.page, query.size = page, page_size
        got = case_list.run(claims, query)
        rows += got.rows
        if not got.has_more:
            break
        page += 1
    columns = [col for col in (product_flow.cfg().get("home_table") or {}).get("columns") or []
               if col.get("key") != "action"]
    out_rows = [[_cell(r, col["key"], lang) for col in columns] for r in rows[:cap]]
    user = str(get_subject(claims) or "")
    return _masked({"title": _say((c.get("title") or {}).get("list"), lang),
                    "subtitle": _say(c.get("generated"), lang, date=_today(), user=user),
                    "sections": [{"heading": "", "columns": [str(language_lock.pick(col.get("label"), lang))
                                                             for col in columns], "rows": out_rows}],
                    "file": _name("cases", filters.get("group") or filters.get("status") or "all")})


def _column(key: str, lang: str) -> str:
    """A home-table column's label (product_flow.yaml home_table.columns) -- one wording everywhere."""
    from app.agents.applicant.copilot.answering import language_lock
    from app.agents.applicant.copilot.capabilities import product_flow

    for col in (product_flow.cfg().get("home_table") or {}).get("columns") or []:
        if col.get("key") == key:
            return str(language_lock.pick(col.get("label"), lang))
    return key


def _cell(row: dict[str, Any], key: str, lang: str) -> str:
    from app.agents.applicant.copilot.capabilities import product_flow

    if key == "app_id":
        return str(row.get("applicant_id") or "")
    if key == "applicant_name":
        return str(row.get("applicant_full_name") or row.get("applicant_name") or "")
    if key == "status":
        return product_flow.status_label(row.get("status") or "", lang)
    if key == "created_date":
        return str(row.get("created_at") or "")[:10]
    return str(row.get(key) or "")


def _name(kind: str, ref: str) -> str:
    raw = str(_cfg().get("file_name") or "{kind}_{ref}_{date}").format(kind=kind, ref=ref, date=_today())
    return "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in raw)[:100]


def _masked(content: dict[str, Any]) -> dict[str, Any]:
    content["title"], content["subtitle"] = _clean(content["title"]), _clean(content["subtitle"])
    for s in content["sections"]:
        s["heading"] = _clean(s["heading"])
        s["columns"] = [_clean(x) for x in s["columns"]]
        s["rows"] = [[_clean(x) for x in row] for row in s["rows"]]
    return content


# --------------------------------------------------------------------------
# the writers
# --------------------------------------------------------------------------

def render(content: dict[str, Any], fmt: str) -> bytes:
    if fmt == "xlsx":
        return _xlsx(content)
    if fmt == "docx":
        return _docx(content)
    if fmt == "pdf":
        return _pdf(content)
    raise ValueError(f"unknown format {fmt}")


def _xlsx(content: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"[:31]
    ws.append([content["title"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([content["subtitle"]])
    for s in content["sections"]:
        ws.append([])
        if s["heading"]:
            ws.append([s["heading"]])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        if s["columns"]:
            ws.append(s["columns"])
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
        for row in s["rows"]:
            ws.append(row)
    for column in ws.columns:
        width = max((len(str(c.value or "")) for c in column), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(10, width + 2), 60)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _markdown(content: dict[str, Any]) -> str:
    lines = [f"# {content['title']}", "", content["subtitle"], ""]
    for s in content["sections"]:
        if s["heading"]:
            lines += [f"## {s['heading']}", ""]
        if s["columns"]:
            lines.append("| " + " | ".join(c.replace("|", "/") for c in s["columns"]) + " |")
            lines.append("|" + "---|" * len(s["columns"]))
            lines += ["| " + " | ".join(str(x).replace("|", "/") for x in row) + " |" for row in s["rows"]]
        else:
            lines += [f"- {row[0]}" for row in s["rows"]]
        lines.append("")
    return "\n".join(lines)


def _pdf(content: dict[str, Any]) -> bytes:
    from app.agents.applicant.copilot.answering import handoff_note

    return handoff_note.to_pdf(_markdown(content))


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _para(text: str, *, bold: bool = False, size: int | None = None) -> str:
    props = ("<w:b/>" if bold else "") + (f'<w:sz w:val="{size * 2}"/>' if size else "")
    run_props = f"<w:rPr>{props}</w:rPr>" if props else ""
    return f'<w:p><w:r>{run_props}<w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'


def _table(columns: list[str], rows: list[list[str]]) -> str:
    def row(cells: list[str], bold: bool) -> str:
        return "<w:tr>" + "".join(f"<w:tc><w:tcPr><w:tcW w:w=\"0\" w:type=\"auto\"/></w:tcPr>{_para(c, bold=bold)}</w:tc>"
                                  for c in cells) + "</w:tr>"

    borders = "".join(f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="999999"/>'
                      for side in ("top", "left", "bottom", "right", "insideH", "insideV"))
    return (f'<w:tbl><w:tblPr><w:tblW w:w="5000" w:type="pct"/><w:tblBorders>{borders}</w:tblBorders></w:tblPr>'
            + (row(columns, True) if columns else "") + "".join(row(r, False) for r in rows) + "</w:tbl>"
            + _para(""))


def _docx(content: dict[str, Any]) -> bytes:
    body = [_para(content["title"], bold=True, size=16), _para(content["subtitle"])]
    for s in content["sections"]:
        if s["heading"]:
            body.append(_para(s["heading"], bold=True, size=13))
        if s["columns"]:
            body.append(_table(s["columns"], s["rows"]))
        else:
            body += [_para("- " + row[0]) for row in s["rows"]]
    document = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="{_W}"><w:body>'
                + "".join(body) + '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/>'
                '<w:pgMar w:top="1000" w:right="1000" w:bottom="1000" w:left="1000"/></w:sectPr></w:body></w:document>')
    files = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        "_rels/.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
        'officeDocument" Target="word/document.xml"/></Relationships>',
        "word/document.xml": document,
    }
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    return buffer.getvalue()


__all__ = ["FORMATS", "allowed", "case_content", "list_content", "render", "reset", "sign", "verify"]
