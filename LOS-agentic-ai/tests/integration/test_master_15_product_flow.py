"""
MASTER SPEC section 15 -- the approved product flow, end to end on the real routes (every dev flag on):

  statuses      stage -> Created / Review / Disbursal -> Pending / Done (config), in the home table and the chat
  home table    columns, search, group / status filter, empty state, row -> the saved form
  follow-up     portfolio answer -> "particular case?" Yes -> "Pending or Done?" -> list -> pick -> open -> "What do
                you want to know?"; No -> the chat goes on; typing anything else skips the flow
  case links    Download Excel / Doc / PDF + Show in UI on case answers; downloads scoped, short-lived, audited
  create / edit chat case creation with the form's rules -> summary -> Confirm; field edit -> Confirm; activity log
  import        row-by-row errors, nothing written until Confirm
  FAQ           "How do I upload a document?" -> numbered steps + source; unknown -> "I don't know that yet"
  login         username + password + stage
"""

from __future__ import annotations

import io
import re
import zipfile

import pytest

from tests.integration.master_env import make_case, prod  # noqa: F401
from tests.integration.test_reupload_supersedes import FOS_SCOPES, _store, client  # noqa: F401

FOS = "/api/v1/fos/copilot"


def turn(client, message, chat="c15", **extra):
    r = client.post(FOS, json={"action": "CUSTOM_QUERY", "message": message, "reply_language": "en",
                               "chat_id": chat, **extra})
    assert r.status_code == 200, r.text
    body = r.json()
    assert set(body) == {"request_id", "markdown", "tts"}, body
    return body


def _href(markdown, label):
    found = re.search(r"\[" + re.escape(label) + r"\]\(((?:action|ask):[^)]*)\)", markdown)
    assert found, (label, markdown)
    return found.group(1)


# ---- statuses + home table ------------------------------------------------------------------------------------
def test_status_mapping_is_config():
    from app.agents.applicant.copilot.capabilities import product_flow as pf

    assert pf.status_of("FOS") == "Created" and pf.status_of("CREDIT") == "Review"
    assert pf.status_of("DISBURSEMENT") == "Disbursal"
    assert pf.group_of("Created") == "pending" and pf.group_of("Review") == "pending"
    assert pf.group_of("Disbursal") == "done"
    assert pf.understand("my pending cases") == {"group": "pending"}
    assert pf.understand("review wale cases") == {"status": "Review"}
    assert pf.understand("disbursal wale") == {"status": "Disbursal"}


def test_home_table_columns_search_filter_and_empty(client, prod):
    empty = client.get("/api/v1/fos/cases").json()
    assert empty["rows"] == [] and empty["empty_text"] == "No cases yet."
    assert [c["key"] for c in empty["columns"]] == ["case_id", "app_id", "applicant_name", "stage", "status",
                                                    "created_date", "action"]
    _, c1 = make_case(client, "Rahul Sharma")
    _, c2 = make_case(client, "Priya Verma")
    page = client.get("/api/v1/fos/cases?group=pending").json()
    assert {r["case_id"] for r in page["rows"]} == {c1, c2}
    row = page["rows"][0]
    assert row["status"] == "Created" and row["group"] == "pending" and row["created_date"]
    assert row["action"]["open_form"].endswith("/form")
    assert client.get("/api/v1/fos/cases?group=done").json()["total"] == 0
    found = client.get("/api/v1/fos/cases?search=Priya").json()
    assert [r["case_id"] for r in found["rows"]] == [c2] and found["rows"][0]["applicant_name"] == "Priya Verma"
    assert client.get("/api/v1/fos/cases?search=Zzzzz").json()["empty_text"] == "No cases match your search."
    assert client.get("/api/v1/fos/cases?group=sideways").status_code == 422


def test_row_click_returns_the_filled_form_and_put_validates(client, prod):
    _, case_id = make_case(client, "Rahul Sharma")
    form = client.get(f"/api/v1/fos/cases/{case_id}/form").json()
    assert form["values"]["full_name"] == "Rahul Sharma" and form["status"] == "Created" and form["editable"]
    bad = client.put(f"/api/v1/fos/cases/{case_id}/form", json={"changes": {"mobile": "12345"}})
    assert bad.status_code == 422 and "mobile" in bad.json()["detail"]["fields"]
    ok = client.put(f"/api/v1/fos/cases/{case_id}/form", json={"changes": {"email": "rahul@example.com"}})
    assert ok.status_code == 200 and ok.json()["values"]["email"] == "rahul@example.com"
    activity = client.get(f"/api/v1/fos/cases/{case_id}/activity").json()["events"]
    assert any(e["type"] == "FORM_UPDATED" and "(ui)" in e["summary"] for e in activity)
    assert any(e["type"] == "CASE_CREATED" for e in activity)
    assert client.get("/api/v1/fos/cases/CASE-DOESNOTEXIST/form").status_code == 404
    schema = client.get("/api/v1/fos/form-schema").json()
    assert [f["name"] for f in schema["fields"]][:2] == ["full_name", "mobile"]


# ---- the chat follow-up ----------------------------------------------------------------------------------------
def test_portfolio_follow_up_yes_pending_pick_open(client, prod):
    make_case(client, "Rahul Sharma")
    make_case(client, "Priya Verma")
    listed = turn(client, "mere cases dikhao")["markdown"]
    assert "Do you want to know about a particular case?" in listed or "Would you like details" in listed
    assert "](ask:Yes)" in listed and "](ask:No)" in listed
    assert "Download: [Excel]" in listed and "Show in UI" in listed         # the list as downloads (one item)
    asked = turn(client, "Yes")["markdown"]
    assert "Pending" in asked and "Done" in asked and "](ask:Pending)" in asked
    pending = turn(client, "Pending")["markdown"]
    assert "Pending" in pending and "Rahul" in pending and "Priya" in pending
    opened = turn(client, "1")["markdown"]
    assert "CASE-" in opened and ("What do you want to know?" in opened or "What would you like to know" in opened)
    assert "](ask:What%20is%20pending?)" in opened
    assert "action:download" not in opened                               # FINAL FIX A4: not on open


def test_portfolio_follow_up_no_and_skip(client, prod):
    make_case(client, "Rahul Sharma")
    turn(client, "my cases", chat="n1")
    said = turn(client, "No", chat="n1")["markdown"]
    assert "Ask me anything" in said or "What else" in said
    turn(client, "my cases", chat="n2")
    other = turn(client, "KYC wale cases", chat="n2")["markdown"]      # skipped: handled as usual
    assert "Pending or Done" not in other


def test_quick_buttons_are_config(client, prod):
    buttons = client.get("/api/v1/fos/quick-actions").json()["buttons"]
    assert [b["label"] for b in buttons] == ["My pending cases", "Create New Case", "How to upload", "FAQ"]
    make_case(client, "Rahul Sharma")
    assert "Rahul" in turn(client, buttons[0]["send"], chat="q1")["markdown"]


# ---- downloads ------------------------------------------------------------------------------------------------
def test_downloads_case_and_list_scoped_and_short_lived(client, prod, make_token):
    _, case_id = make_case(client, "Rahul Sharma")
    files = {}
    for fmt in ("xlsx", "docx", "pdf"):
        link = client.get(f"/api/v1/fos/exports?case_id={case_id}&format={fmt}").json()
        got = client.get(link["url"])
        assert got.status_code == 200 and link["expires_in"] <= 300
        files[fmt] = got.content
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(files["xlsx"])).active
    cells = " ".join(str(c) for row in sheet.iter_rows(values_only=True) for c in row if c)
    assert case_id in cells and "Rahul Sharma" in cells
    assert case_id in zipfile.ZipFile(io.BytesIO(files["docx"])).read("word/document.xml").decode()
    assert files["pdf"].startswith(b"%PDF")
    listed = client.get("/api/v1/fos/exports?format=xlsx&q=group:pending").json()
    assert client.get(listed["url"]).status_code == 200
    # another officer: neither the case nor the link
    from fastapi.testclient import TestClient

    import main

    other = TestClient(main.app)
    other.headers.update({"Authorization": f"Bearer {make_token(subject='other-officer', scopes=FOS_SCOPES)}"})
    assert other.get(f"/api/v1/fos/exports?case_id={case_id}&format=pdf").status_code == 404
    link = client.get(f"/api/v1/fos/exports?case_id={case_id}&format=pdf").json()
    assert other.get(link["url"]).status_code == 403
    assert client.get(f"/api/v1/fos/exports?case_id={case_id}&format=exe").status_code == 422


# ---- create and edit from chat --------------------------------------------------------------------------------
def test_chat_creates_a_case_with_the_form_rules_and_confirm(client, prod):
    started = turn(client, "Create new case", chat="cr")["markdown"]
    assert "Applicant name" in started and "Create New Case" in started
    turn(client, "Sunita Rao", chat="cr")
    bad = turn(client, "12345", chat="cr")["markdown"]
    assert "10-digit" in bad                                             # the form's own rule
    turn(client, "9812345678", chat="cr")
    turn(client, "1988-02-30", chat="cr")                                # not a date: asked again
    turn(client, "1988-02-14", chat="cr")
    turn(client, "45 Park Street, Kolkata", chat="cr")
    turn(client, "personal loan", chat="cr")
    summary = turn(client, "4", chat="cr")["markdown"]                  # 4 rupees: the plausibility rule
    assert "Loan amount" in summary
    summary = turn(client, "450000", chat="cr")["markdown"]
    assert "Confirm?" in summary and "Sunita Rao" in summary and "](action:confirm_write" in summary
    assert client.get("/api/v1/fos/cases?search=Sunita").json()["total"] == 0     # nothing written yet
    done = turn(client, "confirm draft", chat="cr")["markdown"]
    assert re.search(r"CASE-[0-9A-F]+", done) and "created" in done
    assert client.get("/api/v1/fos/cases?search=Sunita").json()["total"] == 1


def test_chat_create_cancel_writes_nothing(client, prod):
    turn(client, "naya case banao", chat="cx")
    turn(client, "Anil Kumar", chat="cx")
    assert "nothing was created" in turn(client, "cancel", chat="cx")["markdown"]
    assert client.get("/api/v1/fos/cases?search=Anil").json()["total"] == 0


def test_chat_edits_a_field_only_after_confirm(client, prod):
    _, case_id = make_case(client, "Rahul Sharma")
    turn(client, f"open {case_id}", chat="ed")
    proposal = turn(client, "loan amount 600000 karo", chat="ed")["markdown"]
    assert "600000" in proposal and "Confirm" in proposal
    assert float(client.get(f"/api/v1/fos/cases/{case_id}/form").json()["values"]["loan_amount"]) == 500000
    confirm = _href(proposal, "Confirm")
    r = client.post(FOS, json={"action_link": confirm, "chat_id": "ed", "reply_language": "en"})
    assert r.status_code == 200 and "Updated" in r.json()["markdown"]
    assert float(client.get(f"/api/v1/fos/cases/{case_id}/form").json()["values"]["loan_amount"]) == 600000
    events = client.get(f"/api/v1/fos/cases/{case_id}/activity").json()["events"]
    assert any(e["type"] == "FORM_UPDATED" and "(chat)" in e["summary"] for e in events)


# ---- import ---------------------------------------------------------------------------------------------------
def test_import_validates_row_by_row_and_writes_only_on_confirm(client, prod):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Applicant name", "Mobile", "Date of birth", "Address", "Product", "Loan amount"])
    ws.append(["Kiran Patil", "9822012345", "1985-06-01", "Pune", "PERSONAL_LOAN", 300000])
    ws.append(["Meena Joshi", "9822098765", "1990-01-20", "Nashik", "PERSONAL_LOAN", 250000])
    ws.append(["Bad Row", "123", "1990-01-20", "Nashik", "PERSONAL_LOAN", 250000])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/v1/fos/imports", files={"file": ("cases.xlsx", buf.getvalue(),
                                                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    check = r.json()
    assert check["ok"] == 2 and check["bad"] == 1 and check["errors"][0].startswith("Row 4: Mobile")
    assert client.get("/api/v1/fos/cases").json()["all_total"] == 0                # nothing written yet
    done = client.post(f"/api/v1/fos/imports/{check['import_id']}/confirm").json()
    assert len(done["created"]) == 2 and done["failed"] == []
    assert client.get("/api/v1/fos/cases").json()["all_total"] == 2
    assert client.post(f"/api/v1/fos/imports/{check['import_id']}/confirm").status_code == 404   # once only
    junk = client.post("/api/v1/fos/imports", files={"file": ("x.xlsx", b"not excel", "application/octet-stream")})
    assert junk.status_code == 422


# ---- FAQ answers, unknown, notifications, login ---------------------------------------------------------------
def test_how_to_upload_is_numbered_steps_with_a_source(client, prod):
    md = turn(client, "How do I upload a document?", chat="f1")["markdown"]
    assert md.startswith("1. ") and "\n2. " in md and "Source:" in md


def test_unknown_faq_says_i_dont_know_yet(client, prod):
    md = turn(client, "How do I apply for a credit card for myself?", chat="f2")["markdown"]
    assert "don't know that yet" in md or "don't have an answer" in md


def test_notifications_and_greeting(client, prod, monkeypatch):
    from app.agents.applicant.copilot.capabilities import case_list

    make_case(client, "Rahul Sharma")
    real = case_list._row
    monkeypatch.setattr(case_list, "_row", lambda a: {**real(a), "days_in_stage": 9})
    case_list.clear_cache()
    items = client.get("/api/v1/fos/notifications").json()["items"]
    assert items and "pending for 9 days" in items[0]["text"]
    assert "pending for 9 days" in turn(client, "hi", chat="g1", new_chat=True)["markdown"].replace("**", "")


def test_login_takes_username_password_and_stage(monkeypatch):
    from fastapi.testclient import TestClient

    import main

    c = TestClient(main.app)
    bad = c.post("/api/v1/auth/login", json={"username": "nobody", "password": "x", "stage": "FOS"})
    assert bad.status_code == 401
    from app.api.routes import auth_api

    monkeypatch.setattr(auth_api, "_authenticate", lambda u, p: True)
    ok = c.post("/api/v1/auth/login", json={"username": "officer", "password": "x", "stage": "fos"})
    if ok.status_code == 500:
        pytest.skip("JWT issuer not configured in this environment")
    assert ok.status_code == 200 and ok.json()["stage"] == "FOS"
    assert c.post("/api/v1/auth/login", json={"username": "officer", "password": "x", "stage": "MOON"}).status_code == 422


def test_chat_import_confirm_creates_the_valid_rows(client, prod):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Mobile", "DOB", "Address", "Product", "Amount"])
    ws.append(["Kiran Patil", "9822012345", "1985-06-01", "Pune", "PERSONAL_LOAN", 300000])
    ws.append(["Bad Row", "123", "1990-01-20", "Nashik", "PERSONAL_LOAN", 250000])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(FOS, files={"file": ("cases.xlsx", buf.getvalue(), "application/octet-stream")},
                    params={"chat_id": "imp"})
    assert r.status_code == 200, r.text
    md = r.json()["markdown"]
    assert "1 of 2 rows are valid" in md and "Row 3: Mobile" in md and "Confirm" in md
    assert client.get("/api/v1/fos/cases").json()["all_total"] == 0
    done = client.post(FOS, json={"action_link": _href(md, "Confirm"), "chat_id": "imp", "reply_language": "en"})
    assert done.status_code == 200 and "1 case(s) created" in done.json()["markdown"]
    assert client.get("/api/v1/fos/cases").json()["all_total"] == 1


def test_open_shows_what_the_form_still_needs(client, prod):
    from app.store import get_repository
    import dataclasses

    _, case_id = make_case(client, "Rahul Sharma")
    repo = get_repository()
    application = repo.get_application(case_id)
    repo.save_application(dataclasses.replace(application, loan_amount=None))
    opened = turn(client, f"open {case_id}", chat="gap")["markdown"]
    assert "This case still needs: Loan amount" in opened


def test_a_case_question_with_no_case_open_asks_which_case_then_answers_it(client, prod):
    make_case(client, "Rahul Sharma")
    make_case(client, "Priya Verma")
    asked = turn(client, "Application status", chat="wc")["markdown"]
    assert asked.startswith("Which case is this about?") and "| CASE-" in asked
    answered = turn(client, "1", chat="wc")["markdown"]
    assert "CASE-" in answered and "Which case" not in answered and "Open a case first" not in answered
    hinglish = turn(client, "status kya hai", chat="wc2")["markdown"]
    assert hinglish.startswith("Which case is this about?") and "don't know that yet" not in hinglish


def test_one_case_is_answered_directly_and_view_all_cases_lists(client, prod):
    make_case(client, "Rahul Sharma")
    assert "Which case" not in turn(client, "Application status", chat="one")["markdown"]
    assert "| CASE-" in turn(client, "view all cases", chat="va")["markdown"]
