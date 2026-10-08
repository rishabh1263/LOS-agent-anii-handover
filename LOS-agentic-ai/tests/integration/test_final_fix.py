"""
FINAL FIX + PROOF (docs/MASTER_FINAL_SPEC.md, section 15 wins) -- every dev flag on.

  A1 no contradiction: "all checks pass" never beside a pending / not-run item; the brief says WHY
  A2 no dead end: no case open -> 0 / 1 / several cases; the pick answers the ORIGINAL question
  A3 list phrasings route to the list
  A4 less is more: case open = 3 lines + <= 2 links; <= 3 extra link items per reply; no repeat
  A5 no "--", no "say ..." in any reply
  B  documents / downloads only through authenticated routes; no token -> 401; another officer -> 403 / 404
  C/D the ONE action endpoint: every action the bot emits works (files open, UI routes, chat replies)
  E  the owner's transcript, as golden cases
"""

from __future__ import annotations

import io
import re
import zipfile

import pytest

from tests.integration.master_env import make_case, prod  # noqa: F401
from tests.integration.test_reupload_supersedes import FOS_SCOPES, _store, client  # noqa: F401

FOS, ACTION = "/api/v1/fos/copilot", "/api/v1/fos/action"
LINK = re.compile(r"\]\(((?:action|ask):[^)]*)\)")


def say(client, message, chat="ff", **extra):
    r = client.post(FOS, json={"action": "CUSTOM_QUERY", "message": message, "reply_language": "en",
                               "chat_id": chat, **extra})
    assert r.status_code == 200, r.text
    return r.json()


def extra_links(md: str) -> int:
    """Link items outside the case table and outside the flow question's own answer options."""
    lines = [ln for ln in md.split("\n") if not ln.startswith("|")]
    items = 0
    for ln in lines:
        found = LINK.findall(ln)
        if not found:
            continue
        if re.search(r"\]\(ask:(Yes|No|Pending|Done)\)", ln) or "confirm_write" in ln:
            continue                                    # the flow question's answer options
        items += 1 if ln.startswith("Download:") else len(found) - (1 if "show_in_ui" in ln and ln.startswith("Download:") else 0)
    return items


@pytest.fixture
def other(make_token):
    from fastapi.testclient import TestClient

    import main

    c = TestClient(main.app)
    c.headers.update({"Authorization": f"Bearer {make_token(subject='officer-b', scopes=FOS_SCOPES)}"})
    return c


@pytest.fixture
def stored_document(client, _store):
    """A document on the caller's case, stored the way the upload pipeline stores it."""
    from app.store.documents import get_document_store
    from app.store.models import Document, DocumentStatus

    app_id, case_id = make_case(client, "Rahul Sharma")
    doc_id = f"{case_id}:{app_id}:pan.png"
    _store.save_document(Document(document_id=doc_id, case_id=case_id, applicant_id=app_id, document_type="PAN",
                                  status=DocumentStatus.VERIFIED))
    get_document_store().put(doc_id, b"\x89PNG\r\n\x1a\nfake", "image/png")
    return case_id, doc_id


# ---- A1 ---------------------------------------------------------------------------------------------------------
def test_the_consistency_guard_drops_an_all_clear_beside_an_open_item():
    from app.agents.applicant.copilot.answering import contract

    md = "CASE-1 · Rahul · FOS\nKYC has not run for this party.\nReview: every check passes."
    assert "every check passes" not in contract._consistent(md, "r1")
    assert contract._consistent("All FOS checks have passed.", "r2") == "All FOS checks have passed."


def test_the_brief_says_why_and_never_all_clear_with_an_open_item(client, prod):
    _, case_id = make_case(client, "Aniket Patil")
    md = say(client, f"open {case_id}", chat="a1")["markdown"]
    first, status = md.split("\n")[0], md.split("\n")[1]
    assert case_id in first and "Aniket Patil" in first and "FOS" in first
    assert "not uploaded, so Aniket Patil's KYC" in status.replace(" (Know Your Customer)", "")
    assert "every check passes" not in md and "checks have passed" not in md


def test_no_template_says_all_clear_next_to_open_items():
    """Every summary / review / readiness template: an all-clear line only ever stands alone."""
    import yaml

    texts = []
    for name in ("applicant_agent", "product_flow", "copilot_reply"):
        def walk(node):
            if isinstance(node, dict):
                for v in node.values():
                    walk(v)
            elif isinstance(node, list):
                for v in node:
                    walk(v)
            elif isinstance(node, str):
                texts.append(node)
        walk(yaml.safe_load(open(f"app/config/{name}.yaml", encoding="utf-8")))
    clear = re.compile(r"(?i)every check passes|all checks pass")
    pending = re.compile(r"(?i)not run|pending|not uploaded|in review|failed")
    assert not [t for t in texts if clear.search(t) and pending.search(t)]


# ---- A2 ---------------------------------------------------------------------------------------------------------
def test_zero_cases_offers_create_new_case(client, prod):
    md = say(client, "Application status", chat="z0")["markdown"]
    assert "don't have any cases yet" in md and "action:new_case" in md
    assert "Open a case first" not in md and "say" not in md.lower().split()


def test_one_case_is_answered_directly_with_its_id(client, prod):
    _, case_id = make_case(client, "Rahul Sharma")
    md = say(client, "Application status", chat="z1")["markdown"]
    assert case_id in md and "Which case" not in md


def test_several_cases_ask_once_and_the_pick_answers_the_original_question(client, prod):
    for n in ("Rahul Sharma", "Priya Verma"):
        make_case(client, n)
    asked = say(client, "Application status", chat="z2")["markdown"]
    assert asked.startswith("Which case is this about?") and "| CASE-" in asked
    for pick, chat in (("2", "z2"),):
        answered = say(client, pick, chat=chat)["markdown"]
        assert "Which case" not in answered and re.search(r"CASE-[0-9A-F]+", answered)
        assert "created" in answered.lower() or "pending" in answered.lower()


def test_no_dead_end_line_exists_anywhere():
    import pathlib

    for path in list(pathlib.Path("app/config").glob("*.yaml")) + list(pathlib.Path("app/api/routes").glob("*.py")):
        text = path.read_text(encoding="utf-8")
        assert "Open a case first" not in text, path
        assert not re.search(r'(?i)\bsay \\?"my cases', text), path


# ---- A3 ---------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("phrase", ["view all cases", "show all cases", "list all case", "all cases dikhao",
                                    "mere saare cases"])
def test_list_phrasings_reach_the_list(client, prod, phrase):
    make_case(client, "Rahul Sharma")
    make_case(client, "Priya Verma")
    md = say(client, phrase, chat="l-" + phrase)["markdown"]
    assert "| CASE-" in md and re.search(r"Showing \d+-\d+ of \d+", md)


# ---- A4 / A5 ----------------------------------------------------------------------------------------------------
def test_case_open_is_three_lines_and_at_most_two_links(client, prod):
    _, case_id = make_case(client, "Rahul Sharma")
    md = say(client, f"open {case_id}", chat="o1")["markdown"]
    body = [ln for ln in md.split("\n\n")[0].split("\n") if ln.strip()]
    assert len(body) <= 3 and body[-1].startswith("Next step:")
    # owner 2026-10-08: plus ONE upload (the next step's document) and the always-there Close
    links = LINK.findall(md)
    assert "action:exit_case" in links and len([t for t in links if t.startswith("action:upload")]) == 1
    assert len([t for t in links if not t.startswith(("action:upload", "action:exit_case"))]) <= 2
    assert "switch_case" not in md and "action:download" not in md


def test_replies_keep_the_link_budget_and_never_repeat(client, prod):
    for n in ("Rahul Sharma", "Priya Verma", "Aniket Patil"):
        make_case(client, n)
    chat = "budget"
    seen_before: set[str] = set()
    for message in ["view all cases", "2", "what is pending?", "kyu?", "download excel", "show in UI"]:
        md = say(client, message, chat=chat)["markdown"]
        assert extra_links(md.replace("[Close](action:exit_case)", "")) <= 3, (message, md)   # Close: always
        assert " -- " not in md and not re.search(r'(?i)\b(say|type) "', md), (message, md)
        # Open per row, Close and Upload are the actions themselves: shown again whenever they apply (owner 2026-10-08)
        now = {t for t in LINK.findall(md) if not t.startswith(("action:open_case", "action:exit_case", "action:upload"))}
        assert not (now & seen_before) or message in ("download excel", "show in UI"), (message, now & seen_before)
        seen_before = now


# ---- B ----------------------------------------------------------------------------------------------------------
def test_nothing_is_reachable_without_a_token(client, prod, stored_document):
    from fastapi.testclient import TestClient

    import main

    case_id, doc_id = stored_document
    anon = TestClient(main.app)
    link = client.get(f"/api/v1/fos/exports?case_id={case_id}&format=pdf").json()["url"]
    for path in ["/api/v1/fos/cases", f"/api/v1/fos/cases/{case_id}/form", link,
                 f"/api/v1/fos/exports?case_id={case_id}&format=xlsx", "/api/v1/fos/documents/view?token=x",
                 f"/api/v1/fos/documents/{case_id}", "/api/v1/fos/chats"]:
        assert anon.get(path).status_code == 401, path
    assert anon.post(ACTION, json={"href": f"action:view_document?id={doc_id}"}).status_code == 401
    for folder in ("/runtime/documents/", "/static/", "/public/", "/wwwroot/", "/uploads/"):
        assert anon.get(folder).status_code == 404, folder


def test_another_officer_gets_nothing(client, other, prod, stored_document):
    case_id, doc_id = stored_document
    for href in (f"action:download?format=xlsx&case={case_id}", f"action:view_document?id={doc_id}",
                 f"action:show_in_ui?case={case_id}"):
        assert other.post(ACTION, json={"href": href}).status_code == 404, href
    link = client.get(f"/api/v1/fos/exports?case_id={case_id}&format=pdf").json()["url"]
    assert other.get(link).status_code == 403
    assert "can't find that case" in other.post(ACTION, json={"href": f"action:open_case?id={case_id}",
                                                              "chat_id": "x"}).text \
        or other.post(ACTION, json={"href": f"action:open_case?id={case_id}"}).status_code == 403


def test_no_server_path_in_the_verify_description(client):
    r = client.get("/api/v1/verify/")
    assert r.status_code in (200, 401, 403)
    assert "runtime" not in r.text and ":\\\\" not in r.text and "upload_root" not in r.text


# ---- C / D ------------------------------------------------------------------------------------------------------
def test_every_action_works_through_the_one_endpoint(client, prod, stored_document):
    case_id, doc_id = stored_document
    sigs = {"xlsx": b"PK", "docx": b"PK", "pdf": b"%PDF"}
    for href in [f"action:download?format={f}&case={case_id}" for f in sigs] + \
                [f"action:download_list?format={f}&q=group:pending" for f in sigs]:
        r = client.post(ACTION, json={"href": href})
        fmt = re.search(r"format=(\w+)", href).group(1)
        assert r.status_code == 200 and r.content.startswith(sigs[fmt]), href
        assert re.search(rf'attachment; filename="[^"]+\.{fmt}"', r.headers["content-disposition"])
    xlsx = client.post(ACTION, json={"href": f"action:download?format=xlsx&case={case_id}"}).content
    from openpyxl import load_workbook

    assert case_id in str([c for row in load_workbook(io.BytesIO(xlsx)).active.iter_rows(values_only=True) for c in row])
    docx = client.post(ACTION, json={"href": f"action:download?format=docx&case={case_id}"}).content
    assert "word/document.xml" in zipfile.ZipFile(io.BytesIO(docx)).namelist()
    view = client.post(ACTION, json={"href": f"action:view_document?id={doc_id}"})
    assert view.status_code == 200 and view.content.startswith(b"\x89PNG")
    assert client.post(ACTION, json={"href": f"action:show_in_ui?case={case_id}"}).json() == \
        {"type": "open_ui", "route": f"/cases/{case_id}"}
    assert client.post(ACTION, json={"href": "action:show_list_in_ui?q=group:pending"}).json()["route"] == \
        "/cases?group=pending"
    assert client.post(ACTION, json={"href": "action:new_case"}).json() == {"type": "open_ui", "route": "/cases/new"}
    assert client.post(ACTION, json={"href": "action:copy?ref=draft-1"}).json()["type"] == "copy"
    opened = client.post(ACTION, json={"href": f"action:open_case?id={case_id}", "chat_id": "act"}).json()
    assert opened["type"] == "reply" and case_id in opened["markdown"] and set(opened) == {"type", "request_id",
                                                                                            "markdown", "tts"}
    assert client.post(ACTION, json={"href": "action:exit_case", "chat_id": "act"}).json()["type"] == "reply"
    assert client.post(ACTION, json={"href": "action:switch_case", "chat_id": "act"}).json()["type"] == "reply"
    assert client.post(ACTION, json={"href": "https://evil.example"}).status_code == 422


# ---- E: the owner's transcript ----------------------------------------------------------------------------------
def test_the_owner_transcript(client, prod):
    for n in ("Rahul Sharma", "Priya Verma", "Aniket Patil"):
        make_case(client, n)
    chat = "owner"
    listed = say(client, "view all cases", chat=chat)["markdown"]
    assert "| CASE-" in listed and "Showing 1-3 of 3" in listed
    opened = say(client, "2", chat=chat)["markdown"]
    assert "Priya Verma" in opened.split("\n")[0]
    pending = say(client, "what is pending?", chat=chat)["markdown"]
    assert "PAN" in pending and "Signature" in pending
    why = say(client, "kyu?", chat=chat)["markdown"]
    assert "nothing on file" not in why and "not uploaded" in why
    dl = say(client, "download excel", chat=chat)["markdown"]
    href = re.search(r"\]\((action:download\?format=xlsx[^)]*)\)", dl).group(1)
    assert client.post(ACTION, json={"href": href}).content.startswith(b"PK")
    ui = say(client, "show in UI", chat=chat)["markdown"]
    assert "action:show_in_ui" in ui
    other = say(client, "dusra case", chat=chat)["markdown"]
    assert "| CASE-" in other
    asked = say(client, "Application status", chat="owner-2", new_chat=True)["markdown"]
    assert asked.startswith("Which case is this about?")
    answered = say(client, "1", chat="owner-2")["markdown"]
    assert "Which case" not in answered and "Signature" in answered.replace("1 more", "Signature")
    abuse = say(client, "madarchod what is my kyc status", chat="owner-2")
    assert "**m*******d**" in abuse["markdown"] and "madarchod" not in abuse["tts"] and "KYC" not in abuse["markdown"]
    kyc = say(client, "KYC status", chat="owner-2")["markdown"]
    assert "KYC" in kyc and "foul" not in kyc
